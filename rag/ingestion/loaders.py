"""Document loaders: turn raw bytes of PDF, DOCX, HTML and Markdown into ``ParsedDocument`` sections."""

from __future__ import annotations

import io
import re
from collections.abc import Callable
from pathlib import PurePath

from rag.models import ParsedDocument, Section

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")


class UnsupportedFileType(ValueError):
    pass


class DocumentParseError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def split_markdown_sections(text: str) -> list[Section]:
    """Split markdown into sections at headings. ``heading`` is the full path, e.g. ``Guide > Setup``.

    Headings inside fenced code blocks are ignored.
    """
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []
    in_fence = False

    def flush() -> None:
        body = "\n".join(buf).strip()
        heading = " > ".join(h for _, h in stack) or None
        if body:
            sections.append(Section(text=body, heading=heading))
        buf.clear()

    for line in text.splitlines():
        if _FENCE_RE.match(line):
            in_fence = not in_fence
            buf.append(line)
            continue
        m = None if in_fence else _HEADING_RE.match(line)
        if m:
            flush()
            level = len(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, m.group(2).strip()))
        else:
            buf.append(line)
    flush()
    return sections


def _first_heading(sections: list[Section]) -> str | None:
    for s in sections:
        if s.heading:
            return s.heading.split(" > ")[0]
    return None


def load_markdown(name: str, data: bytes) -> ParsedDocument:
    text = _decode(data)
    sections = split_markdown_sections(text)
    return ParsedDocument(name=name, title=_first_heading(sections), sections=sections)


def load_text(name: str, data: bytes) -> ParsedDocument:
    text = _decode(data).strip()
    return ParsedDocument(name=name, sections=[Section(text=text)] if text else [])


def _unwrap_pdf_lines(text: str) -> str:
    """PDF text comes back hard-wrapped; re-join lines that do not end a sentence into paragraphs."""
    text = re.sub(r"[ \t]+\n", "\n", text.strip())
    text = re.sub(r"(?<=[^.!?:\n])\n(?!\n)", " ", text)
    return re.sub(r"\n(?!\n)", "\n\n", text)


def load_pdf(name: str, data: bytes) -> ParsedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentParseError(f"{name}: encrypted PDFs are not supported")
        sections = []
        for i, page in enumerate(reader.pages, start=1):
            text = _unwrap_pdf_lines(page.extract_text() or "")
            if text:
                sections.append(Section(text=text, page=i))
        title = None
        if reader.metadata and reader.metadata.title:
            title = str(reader.metadata.title)
        return ParsedDocument(
            name=name, title=title, sections=sections, metadata={"pages": len(reader.pages)}
        )
    except PdfReadError as exc:
        raise DocumentParseError(f"{name}: could not read PDF ({exc})") from exc


def load_docx(name: str, data: bytes) -> ParsedDocument:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated exception types
        raise DocumentParseError(f"{name}: could not read DOCX ({exc})") from exc

    lines: list[str] = []
    # When the document uses a "Title" paragraph, nest "Heading N" one level below it.
    offset = 1 if any(p.style is not None and p.style.name == "Title" for p in document.paragraphs) else 0
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        if style == "Title":
            lines.append(f"# {text}")
        elif style.startswith("Heading"):
            try:
                level = int(style.split()[-1])
            except ValueError:
                level = 1
            lines.append(f"{'#' * min(max(level, 1) + offset, 6)} {text}")
        elif style.startswith("List"):
            lines.append(f"- {text}")
        else:
            lines.append(text)
        lines.append("")
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
        lines.append("")

    sections = split_markdown_sections("\n".join(lines))
    title = document.core_properties.title or _first_heading(sections)
    return ParsedDocument(name=name, title=title or None, sections=sections)


_BLOCK_TAGS = [
    "p",
    "div",
    "section",
    "article",
    "li",
    "ul",
    "ol",
    "tr",
    "table",
    "pre",
    "blockquote",
    "dd",
    "dt",
    "header",
    "main",
    "aside",
    "figcaption",
    "br",
]


def load_html(name: str, data: bytes) -> ParsedDocument:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(_decode(data), "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    for tag in soup(["script", "style", "noscript", "template", "svg", "nav", "footer", "form"]):
        tag.decompose()
    body = soup.body or soup
    # HTML whitespace (including newlines) is insignificant, so mark real breaks with a sentinel,
    # collapse all whitespace, then turn the sentinels back into newlines.
    brk = "\u2029"
    for level in range(1, 7):
        for h in body.find_all(f"h{level}"):
            text = " ".join(h.get_text(" ", strip=True).split())
            h.replace_with(f"{brk}{brk}{'#' * level} {text}{brk}{brk}")
    for tag in body.find_all("li"):
        tag.insert(0, "- ")
    for tag in body.find_all(["td", "th"]):
        tag.insert_after(" | ")
    for tag in body.find_all(_BLOCK_TAGS):
        tag.insert_before(brk)
        tag.insert_after(brk)

    raw = re.sub(r"[^\S\u2029]+", " ", body.get_text())
    lines = [line.strip() for line in raw.split(brk)]
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    text = re.sub(r"[ \t]*\|[ \t]*$", "", text, flags=re.MULTILINE)
    sections = split_markdown_sections(text)
    return ParsedDocument(name=name, title=title or _first_heading(sections), sections=sections)


LOADERS: dict[str, Callable[[str, bytes], ParsedDocument]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".html": load_html,
    ".htm": load_html,
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".txt": load_text,
}
SUPPORTED_EXTENSIONS = tuple(sorted(LOADERS))


def load_document(name: str, data: bytes) -> ParsedDocument:
    """Parse ``data`` using the loader matching the file extension of ``name``."""
    ext = PurePath(name).suffix.lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise UnsupportedFileType(
            f"Unsupported file type '{ext or name}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    return loader(name, data)
