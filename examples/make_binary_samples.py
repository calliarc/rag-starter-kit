"""Regenerate the binary sample documents (PDF and DOCX) in examples/docs.

    python examples/make_binary_samples.py

The PDF is written with a tiny built-in writer (no extra dependency); the DOCX uses python-docx.
All content is fictional sample text for the RAG Starter Kit demo.
"""

from __future__ import annotations

import textwrap
from pathlib import Path

DOCS = Path(__file__).parent / "docs"

SECURITY_POLICY_PAGES = [
    [
        "Fernhill Instruments Information Security Policy",
        "",
        "Sample content for the RAG Starter Kit demo. Fernhill Instruments is a fictional company.",
        "",
        "1. Passwords and authentication",
        "All accounts must use a password of at least 14 characters. Passwords are stored in the company "
        "password manager and must never be reused across services. Multi-factor authentication is "
        "mandatory for email, source code hosting, the cloud console and the Fernhill Portal. Hardware "
        "security keys are required for administrators.",
        "",
        "2. Access reviews",
        "Team leads review who has access to production systems every quarter. Access for people who "
        "leave the company is removed on their last working day.",
        "",
        "3. Laptops and devices",
        "Company laptops use full-disk encryption and lock automatically after five minutes of inactivity. "
        "Personal devices may only access email through the managed mobile app.",
    ],
    [
        "4. Incident response",
        "Report any suspected security incident to the security team within one hour of discovery, using "
        "the #security-incidents channel or the on-call phone number in the IT portal. Do not try to "
        "investigate on your own and do not delete evidence.",
        "",
        "Incidents are classified by severity. Severity 1 means customer data is exposed or production is "
        "down; the incident commander is paged immediately and customers are informed within 72 hours "
        "when their data is affected. Severity 2 covers contained incidents without customer impact.",
        "",
        "5. Data retention",
        "Security logs are retained for 400 days. Customer sensor data is kept for five years unless the "
        "customer requests earlier deletion. Backups are encrypted and tested for restore every month.",
        "",
        "6. Policy owner",
        "This policy is owned by the Head of Security and reviewed once per year.",
    ],
]


def _pdf_escape(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def write_simple_pdf(path: Path, pages: list[list[str]], title: str = "") -> None:
    """Write a text-only PDF (Helvetica, A4). Each item in ``pages`` is a list of paragraphs."""
    objects: list[bytes] = []

    def add(obj: bytes) -> int:
        objects.append(obj)
        return len(objects)

    font_id = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    pages_id = add(b"")  # placeholder, filled below
    page_ids = []
    for paragraphs in pages:
        lines: list[str] = []
        for para in paragraphs:
            lines.extend(textwrap.wrap(para, 92) or [""])
        ops = ["BT", "/F1 11 Tf", "14 TL", "56 790 Td"]
        for line in lines:
            ops.append(f"({_pdf_escape(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        content_id = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                (
                    f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 595 842] "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
                ).encode()
            )
        )
    kids = " ".join(f"{i} 0 R" for i in page_ids)
    objects[pages_id - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()
    catalog_id = add(f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode())
    info_id = add(f"<< /Title ({_pdf_escape(title)}) /Producer (rag-starter-kit) >>".encode())

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root {catalog_id} 0 R /Info {info_id} 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    path.write_bytes(bytes(out))


def write_onboarding_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.title = "New Hire Onboarding Guide"
    d.add_heading("New Hire Onboarding Guide", level=0)
    d.add_paragraph(
        "Sample content for the RAG Starter Kit demo. Fernhill Instruments is a fictional company."
    )
    d.add_heading("Before your first day", level=1)
    d.add_paragraph(
        "About a week before you start, People Operations sends you a welcome email with your start time, "
        "the office address and a link to sign your contract electronically."
    )
    d.add_heading("Your first day", level=1)
    d.add_paragraph(
        "Your first day starts at 09:30 at the reception desk, where your manager will meet you. The IT "
        "help desk hands over your laptop and helps you set up multi-factor authentication before lunch."
    )
    d.add_paragraph(
        "In the afternoon you will join a welcome session that introduces our products and customers."
    )
    d.add_heading("Onboarding buddy", level=1)
    d.add_paragraph(
        "Every new hire is paired with an onboarding buddy from another team. Your buddy is your go-to "
        "person for informal questions during your first 90 days, and you will meet at least once a week."
    )
    d.add_heading("Required training", level=1)
    d.add_paragraph(
        "Complete the security awareness and data protection courses within your first 14 days. Lab staff "
        "must also finish the chemical safety module before working at a calibration bench."
    )
    for item in (
        "Security awareness (45 minutes)",
        "Data protection basics (30 minutes)",
        "Chemical safety for lab staff (60 minutes)",
    ):
        d.add_paragraph(item, style="List Bullet")
    d.add_heading("Probation period", level=1)
    d.add_paragraph(
        "The probation period is six months. Your manager will hold check-ins after 30, 90 and 150 days."
    )
    d.save(str(path))


if __name__ == "__main__":
    DOCS.mkdir(parents=True, exist_ok=True)
    write_simple_pdf(DOCS / "security-policy.pdf", SECURITY_POLICY_PAGES, title="Information Security Policy")
    write_onboarding_docx(DOCS / "onboarding-guide.docx")
    print("wrote", DOCS / "security-policy.pdf", "and", DOCS / "onboarding-guide.docx")
